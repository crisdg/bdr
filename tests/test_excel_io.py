import openpyxl
from helpers import make_row

from bdr_leader_merge.compare import compare_files
from bdr_leader_merge.excel_io import read_rows, write_merged
from bdr_leader_merge.merge import merge_estimados


def test_write_merged_conserva_las_demas_columnas(tmp_path, build_wb):
    p_n = build_wb(
        "n.xlsx",
        [
            make_row("000015", "09", "000020", marca="SEX FOR WOMEN"),
            make_row("000015", "14", "000000", marca="SEX FOR WOMEN"),
        ],
    )
    p_n1 = build_wb(
        "n1.xlsx",
        [
            make_row("000015", "09", "000020", marca="SEX FOR WOMEN"),
            make_row("000999", "01", "000200", marca="NUEVO"),
        ],
    )
    out = tmp_path / "out.xlsx"

    res = merge_estimados(read_rows(p_n), read_rows(p_n1))
    write_merged(p_n, p_n1, out, res)

    filas = read_rows(out)
    assert len(filas) == 3

    assert filas[0][40] == 21
    assert filas[1][40] == 0
    assert filas[2][40] == 10

    # columnas ajenas al estimado intactas en las filas de n
    originales = read_rows(p_n)
    for i in range(2):
        for c in range(49):
            if c != 40:
                assert filas[i][c] == originales[i][c]

    # la fila nueva es copia de n+1 salvo el estimado
    fuente = read_rows(p_n1)[1]
    for c in range(49):
        if c != 40:
            assert filas[2][c] == fuente[c]
    assert filas[2][48] == "NUEVO"


def test_estimado_se_escribe_numerico(tmp_path, build_wb):
    p_n = build_wb("n.xlsx", [make_row("000015", "09", "000020")])
    p_n1 = build_wb("n1.xlsx", [make_row("000015", "09", "000020")])
    out = tmp_path / "out.xlsx"

    res = merge_estimados(read_rows(p_n), read_rows(p_n1))
    write_merged(p_n, p_n1, out, res)

    wb = openpyxl.load_workbook(out)
    celda = wb[wb.sheetnames[0]].cell(row=2, column=41)
    assert isinstance(celda.value, int)
    assert celda.value == 21
    wb.close()


def test_compare_detecta_diferencias(tmp_path, build_wb):
    p_n = build_wb("n.xlsx", [make_row("000015", "09", "000020")])
    p_n1 = build_wb("n1.xlsx", [make_row("000015", "09", "000020")])
    out = tmp_path / "out.xlsx"
    res = merge_estimados(read_rows(p_n), read_rows(p_n1))
    write_merged(p_n, p_n1, out, res)

    igual = build_wb("ref_ok.xlsx", [make_row("000015", "09", 21)])
    rep = compare_files(out, igual)
    assert rep.diferencias_estimado == []
    assert rep.diferencias_otras == []
    assert rep.celdas_estimado_iguales == 1

    distinto = build_wb("ref_dif.xlsx", [make_row("000015", "09", 99)])
    rep = compare_files(out, distinto)
    assert len(rep.diferencias_estimado) == 1
    assert rep.diferencias_estimado[0].calculado == 21
    assert rep.diferencias_estimado[0].referencia == 99
