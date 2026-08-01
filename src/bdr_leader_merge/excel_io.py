"""Lectura y escritura de las planillas LEADER preservando el archivo base."""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .config import DEFAULT_CONFIG, MergeConfig
from .merge import MergeResult
from .model import is_blank_row


def read_rows(path: Path, cfg: MergeConfig = DEFAULT_CONFIG) -> list[list[Any]]:
    """Devuelve las filas de datos de la primera hoja, sin encabezado ni filas vacias."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        filas = [list(r) for r in ws.iter_rows(min_row=cfg.first_data_row, values_only=True)]
    finally:
        wb.close()
    return [f for f in filas if not is_blank_row(f)]


def _last_data_row(ws: Worksheet, cfg: MergeConfig) -> int:
    """max_row de openpyxl incluye filas con solo formato; se busca la ultima con datos."""
    fila = ws.max_row
    while fila >= cfg.first_data_row:
        valores = [ws.cell(row=fila, column=c).value for c in range(1, ws.max_column + 1)]
        if not is_blank_row(valores):
            return fila
        fila -= 1
    return cfg.first_data_row - 1


def write_merged(
    path_n: Path,
    path_n1: Path,
    path_out: Path,
    result: MergeResult,
    cfg: MergeConfig = DEFAULT_CONFIG,
) -> Path:
    """Escribe el resultado partiendo del archivo base para no alterar el resto de columnas."""
    wb_out = openpyxl.load_workbook(path_n)
    ws_out = wb_out[wb_out.sheetnames[0]]
    wb_src = openpyxl.load_workbook(path_n1)
    ws_src = wb_src[wb_src.sheetnames[0]]

    try:
        ultima = _last_data_row(ws_out, cfg)

        for idx, estimado in result.estimados_base.items():
            celda = ws_out.cell(row=idx + cfg.first_data_row, column=cfg.col_estimado)
            celda.value = estimado
            celda.number_format = "General"

        destino = ultima + 1
        n_cols = ws_src.max_column
        for src_idx, estimado in result.filas_nuevas:
            fila_src = src_idx + cfg.first_data_row
            for col in range(1, n_cols + 1):
                origen = ws_src.cell(row=fila_src, column=col)
                celda = ws_out.cell(row=destino, column=col)
                celda.value = origen.value
                celda.number_format = origen.number_format
                celda.alignment = copy(origen.alignment)
            celda_ao = ws_out.cell(row=destino, column=cfg.col_estimado)
            celda_ao.value = estimado
            celda_ao.number_format = "General"
            destino += 1

        path_out.parent.mkdir(parents=True, exist_ok=True)
        wb_out.save(path_out)
    finally:
        wb_out.close()
        wb_src.close()

    return path_out
